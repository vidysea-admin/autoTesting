"""Tester-style run reports: an Excel summary and a screen-by-screen HTML
report with embedded screenshots. Contract: qa/contracts/report-export.md.

Both read only what `ProjectStore` already has (cases, RawResults, Verdicts)
— exporting is presentation over existing evidence, never a new source of
truth (design principle 8, same discipline as `ui/`).
"""

from __future__ import annotations

import base64
from datetime import UTC
from html import escape
from pathlib import Path

from openpyxl import Workbook

from autotester.browser.secrets import SecretStore
from autotester.core.excel import append_row, autosize_columns
from autotester.core.redact import Redactor
from autotester.schema.case import Case
from autotester.schema.enums import EvidenceKind, Result
from autotester.schema.run import Run
from autotester.schema.verdict import Verdict
from autotester.stages.run_artifacts import bundle_report_detail
from autotester.stages.ux_advisory import add_ux_sheet, load_ux_state, ux_findings_html
from autotester.store import ProjectStore

_BADGE_COLOR = {
    "PASS": "#16a34a", "FAIL": "#dc2626", "BLOCKED": "#b45309", "INCONCLUSIVE": "#6b7280",
}


def _run_sort_key(run: Run) -> tuple:
    created = run.created_at if run.created_at.tzinfo else run.created_at.replace(tzinfo=UTC)
    return created.astimezone(UTC), run.id


def valid_runs_newest_first(store: ProjectStore) -> list[Run]:
    """Persisted Run envelopes only; auxiliary run artifacts are not runs."""
    runs_dir = store.paths.runs_dir
    if not runs_dir.exists():
        return []
    runs: list[Run] = []
    for path in sorted((p for p in runs_dir.iterdir() if p.is_dir()), reverse=True):
        try:
            run = store.load_run(path.name)
        except ValueError:
            continue
        if run is not None and run.id == path.name and run.project == store.paths.slug:
            runs.append(run)
    runs.sort(key=_run_sort_key, reverse=True)
    return runs


def bounds_text(run: Run | None) -> str:
    """AT-682/CN10: the bounds a run ran under, raw (never capped); None = NOT RECORDED."""
    b = run.bounds if run else None
    if b is None:
        return "approval bounds NOT RECORDED"
    return (f"ran under approval {b.approval_id}: {b.max_actions} actions / "
            f"{b.max_probes} probes / {b.wall_clock_s}s")


def _latest_run_id(store: ProjectStore) -> str:
    runs = valid_runs_newest_first(store)
    if not runs:
        raise ValueError(f"no runs exist yet for '{store.paths.slug}'")
    return runs[0].id


def _case_lookup(store: ProjectStore) -> dict[str, Case]:
    return {c.id: c for c in store.list_cases()}


def _load_redactor(store: ProjectStore) -> Redactor:
    """C5/AT-594: load the shared last-stop Redactor, including for repro steps."""
    project = store.load_project()
    if project is None:
        return Redactor({})
    return SecretStore.load(project, store.paths.env_file, strict=False).redactor()


def _needs_developer_detail(verdict: Verdict | None) -> bool:
    """RE6 (D-045): a FAIL or INCONCLUSIVE verdict is what a developer must act
    on -- a PASS gets no failure/repro detail."""
    return verdict is not None and verdict.result in (Result.FAIL, Result.INCONCLUSIVE)


def _failure_rows(verdict: Verdict | None) -> list[tuple[str, str, str | None]]:
    """Each failure's (criterion_id, reason, fix_hint), verbatim off the
    stored Verdict -- RE1: nothing recomputed."""
    if verdict is None:
        return []
    return [(f.criterion_id, f.reason, f.fix_hint) for f in verdict.failures]


def _repro_steps(case: Case | None, redactor: Redactor) -> list[str]:
    """AT-594: format the case's own steps, preserving placeholders and scrubbing values."""
    if case is None:
        return []
    return [
        redactor.scrub(
            f"{step.order}. {step.action.value} {step.target}"
            + (f" = {step.value}" if step.value else "")
        )
        for step in case.steps
    ]


def export_excel(
    project_slug: str, run_id: str | None, out_path: Path, root: Path | None = None,
    *, case_overrides: dict[str, Case] | None = None,
) -> Path:
    """One row per case: what ran, what it was checked against, what happened."""
    store = ProjectStore(project_slug, root)
    run_id = run_id or _latest_run_id(store)
    cases = _case_lookup(store) if case_overrides is None else case_overrides
    verdicts = {v.case_id: v for v in store.load_verdicts(run_id)}
    redactor = _load_redactor(store)
    ux = load_ux_state(store, run_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Run report"
    ws.append(["Case", "Kind", "Class", "Outcome", "Result", "Criteria met",
               "Duration (s)", "Grader", "Notes", "Failures", "Repro steps", "Failure bundle"])
    for result in store.load_results(run_id):
        case = cases.get(result.case_id)
        verdict = verdicts.get(result.case_id)
        detail = _needs_developer_detail(verdict)
        failures_text = "\n".join(
            f"{criterion_id}: {reason}" + (f" (fix: {fix_hint})" if fix_hint else "")
            for criterion_id, reason, fix_hint in _failure_rows(verdict)
        ) if detail else ""
        steps_text = "\n".join(_repro_steps(case, redactor)) if detail else ""
        case_title = redactor.scrub(case.title) if case else result.case_id
        append_row(ws, [
            case_title,
            case.kind.value if case else "",
            case.case_class.value if case else "",
            result.outcome.value,
            verdict.result.value if verdict else "",
            f"{verdict.criteria_met}/{verdict.criteria_total}" if verdict else "",
            round(result.duration_s, 2),
            verdict.grader_provider if verdict else "",
            (verdict.scoreboard if verdict else "") or (result.error or ""),
            failures_text,
            steps_text,
        ])
        href, note = bundle_report_detail(store, run_id, result.case_id, out_path.parent)
        cell = ws.cell(ws.max_row, 12, note)
        if href is not None:
            cell.hyperlink = href
            cell.style = "Hyperlink"
    autosize_columns(ws)
    add_ux_sheet(wb, ux, cases, redactor)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def png_base64(path: Path, allowed_root: Path, trusted_root: Path) -> str | None:
    """Base64-encode a screenshot for inline embedding. Public — also used by
    `ui/routes_report.py` to embed the same screenshots in the live view, so
    the live page and the portable export show identical evidence."""
    trusted = trusted_root.resolve()
    try:
        parts = allowed_root.absolute().relative_to(trusted).parts
    except ValueError:
        return None
    safe_root = trusted
    for part in parts:
        safe_root /= part
        if safe_root.is_symlink():
            return None
    resolved_safe_root = safe_root.resolve()
    if resolved_safe_root != safe_root:
        return None
    candidate = path.resolve()
    if not candidate.is_relative_to(resolved_safe_root):
        return None
    if not candidate.is_file():
        return None
    return base64.b64encode(candidate.read_bytes()).decode("ascii")


def _case_section(
    store: ProjectStore, run_id: str, case: Case | None, result, verdict, redactor: Redactor,
    ux=(None, None),
    export_dir: Path | None = None,
) -> str:
    run_dir = store.paths.run_dir(run_id)
    # AT-609: `Case.title` and `Evidence.label` are operator free text: scrubbed here, then escaped.
    note = escape(f" [{case.provenance}]") if case and case.provenance else ""  # D-097
    title = escape(redactor.scrub(case.title) if case else result.case_id)
    color = _BADGE_COLOR.get(verdict.result.value if verdict else "", "#6b7280")
    badge_text = escape(verdict.result.value) if verdict else escape(result.outcome.value)
    shots = [e for e in result.evidence if e.kind is EvidenceKind.SCREENSHOT]
    figures = "".join(
        f"<figure><img src='data:image/png;base64,{data}'>"
        f"<figcaption>{escape(redactor.scrub(shot.label or shot.path))}</figcaption></figure>"
        for shot in shots
        for data in [png_base64(run_dir / shot.path, run_dir, store.paths.dir)] if data is not None
    )
    scoreboard = escape(verdict.scoreboard) if verdict and verdict.scoreboard else ""
    error = escape(result.error) if result.error else ""
    no_shots = "<p class='meta'>no screenshots captured</p>"
    detail = _failure_detail_html(verdict, case, redactor)
    href, bundle_note = bundle_report_detail(store, run_id, result.case_id, export_dir)
    bundle_html = (f"<a href='{escape(href, quote=True)}'>{escape(bundle_note)}</a>"
                   if href else escape(bundle_note))
    return (
        f"<section><h2>{title} "
        f"<span class='badge' style='background:{color}'>{badge_text}</span></h2>"
        f"<p class='meta'>{scoreboard}{error}{note}</p>"
        f"{detail}{ux_findings_html(ux, result.case_id, redactor)}"
        f"<div class='shots'>{figures or no_shots}</div>"
        f"{_video_link_html(result)}<p class='meta'>{bundle_html}</p>"
        "</section>"
    )


def _video_link_html(result) -> str:
    """RE3's named exception (D-050) + T-191 V8: a kept video is LINKED, never
    embedded like a screenshot -- a 15-20 minute recording base64'd into the
    page would defeat the whole point of the exception. Path is run-relative,
    same as `Evidence.path` everywhere else; `report-export.md`'s "relative
    to the run directory" phrasing (not to the exported file)."""
    return "".join(
        f"<p class='meta'><a href='{escape(v.path)}'>video ({escape(v.path)})</a></p>"
        for v in result.evidence if v.kind is EvidenceKind.VIDEO
    )


def _failure_detail_html(verdict: Verdict | None, case: Case | None, redactor: Redactor) -> str:
    """RE6 (D-045): for a FAIL/INCONCLUSIVE verdict, each failure's criterion,
    reason and fix_hint, plus the case's own steps as the repro -- read
    straight off the stored Verdict/Case (RE1: nothing recomputed)."""
    if not _needs_developer_detail(verdict):
        return ""
    items = "".join(
        f"<li><code>{escape(criterion_id)}</code> — {escape(reason)}"
        + (f" <em>fix: {escape(fix_hint)}</em>" if fix_hint else "") + "</li>"
        for criterion_id, reason, fix_hint in _failure_rows(verdict)
    )
    failures_html = f"<h3>Failures</h3><ul>{items}</ul>" if items else ""
    steps = _repro_steps(case, redactor)
    steps_html = (
        "<h3>Repro steps</h3><ol>"
        + "".join(f"<li>{escape(step)}</li>" for step in steps)
        + "</ol>"
    ) if steps else ""
    both = failures_html + steps_html
    return f"<div class='detail'>{both}</div>" if both else ""


def export_html(
    project_slug: str, run_id: str | None, out_path: Path, root: Path | None = None,
    *, case_overrides: dict[str, Case] | None = None,
) -> Path:
    """One section per case, in run order, each with its own screenshots
    embedded inline (base64) so the file is a single portable artifact."""
    store = ProjectStore(project_slug, root)
    run_id = run_id or _latest_run_id(store)
    cases = _case_lookup(store) if case_overrides is None else case_overrides
    verdicts = {v.case_id: v for v in store.load_verdicts(run_id)}
    results = store.load_results(run_id)
    redactor = _load_redactor(store)
    ux = load_ux_state(store, run_id)

    sections = "".join(
        _case_section(store, run_id, cases.get(r.case_id), r, verdicts.get(r.case_id), redactor, ux,
                      out_path.parent)
        for r in results
    )
    style = (
        "body{font-family:-apple-system,Segoe UI,Roboto,sans-serif;max-width:900px;"
        "margin:2rem auto;padding:0 1rem;color:#16181d}"
        "h1{margin-bottom:.2rem}.meta{color:#667085;font-size:.85rem}"
        "section{border:1px solid #e2e5ea;border-radius:10px;padding:1.2rem 1.4rem;"
        "margin-bottom:1.2rem}"
        ".badge{color:#fff;padding:.15rem .6rem;border-radius:999px;font-size:.75rem}"
        ".detail{margin-top:.6rem}.detail h3{font-size:.85rem;margin:.6rem 0 .2rem}"
        ".detail li{font-size:.85rem;margin-bottom:.2rem}"
        ".shots{display:flex;flex-wrap:wrap;gap:1rem;margin-top:.8rem}"
        "figure{margin:0;max-width:320px}img{max-width:100%;border:1px solid #e2e5ea;"
        "border-radius:6px}figcaption{font-size:.75rem;color:#667085;margin-top:.3rem}"
    )
    html = (
        "<!doctype html><meta charset='utf-8'>"
        f"<title>{escape(project_slug)} — {escape(run_id)} report</title>"
        f"<style>{style}</style>"
        f"<h1>{escape(project_slug)} — run report</h1>"
        f"<p class='meta'>Run <code>{escape(run_id)}</code> · {len(results)} case(s)</p>"
        f"<p class='meta'>{escape(bounds_text(store.load_run(run_id)))}</p>"
        f"{sections}"
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(html, encoding="utf-8")
    return out_path
