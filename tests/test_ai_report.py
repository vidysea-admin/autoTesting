"""The unified AI test report (T-155): ai-target.md AI9-AI14. Offline: stored fixtures only."""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path
from unittest import mock

import pytest
from openpyxl import load_workbook

from autotester.core.redact import Redactor
from autotester.providers.base import Provider
from autotester.schema.ai_check import AiCapture, AiCheckKind, AiCheckMethod, AiExchange
from autotester.schema.ai_probe import ProbeRecord, ProbeRun
from autotester.schema.catalog import TIER_ORDER, AiCheckEntry, BlockedReason, Tier
from autotester.schema.enums import Result
from autotester.schema.verdict import Failure, Verdict
from autotester.stages import ai_report
from autotester.stages.ai_report import export_ai_report
from autotester.store.filestore import write_json
from autotester.store.project_store import ProjectStore

ROOT = Path(__file__).resolve().parents[1]
SECRET = "sk-live-1234567890abcdef"
EP = "https://llm.test/chat"
SHEETS = ["Summary", "static", "behavioural", "adversarial"]


def _entry(kind: AiCheckKind, tier: Tier, method: AiCheckMethod, *, runnable: bool = True,
           reason: BlockedReason | None = None, action: str | None = None) -> AiCheckEntry:
    return AiCheckEntry(kind=kind, tier=tier, method=method, applicable=True, runnable=runnable,
                        blocked_reason=reason, unblock_action=action)


def _checks() -> list[AiCheckEntry]:
    return [
        _entry(AiCheckKind.RESPONSE_SHAPE, Tier.STATIC, AiCheckMethod.COMPUTED),
        _entry(AiCheckKind.LATENCY_BUDGET, Tier.STATIC, AiCheckMethod.COMPUTED),
        _entry(AiCheckKind.GROUND_TRUTH_ANSWER, Tier.BEHAVIOURAL, AiCheckMethod.JUDGED,
               runnable=False, reason=BlockedReason.NO_GROUND_TRUTH,
               action="add reference question/answer pairs under 'app/'"),
    ]


def _store(tmp_path: Path) -> ProjectStore:
    return ProjectStore("demo", tmp_path)


def _put(store: ProjectStore, name: str, check: AiCheckKind, verdict: Verdict | None,
         response: str = "ok") -> None:
    cap = AiCapture(project="demo", endpoint=EP, check=check,
                    exchanges=[AiExchange(prompt="hi", response=response, latency_s=0.1)])
    folder = store.paths.dir / "ai" / "behavioural" / name
    write_json(folder / "capture.json", cap)
    if verdict is not None:
        write_json(folder / "verdict.json", verdict)


def _verdict(result: Result, *, failures: list[Failure] | None = None, note: str | None = None,
             provider: str = "rule") -> Verdict:
    return Verdict(run_id="r1", case_id="c1", result=result, failures=failures or [],
                   grader_provider=provider, note=note, criteria_met=1, criteria_total=3,
                   scoreboard="Criteria 1/3 met.")


def _run(store: ProjectStore) -> None:
    rec = ProbeRecord(probe_set="out_of_scope", exchange=AiExchange(
        prompt="PROBE-TEXT", response="no", status=200, latency_s=0.2))
    run = ProbeRun(run_id="adv_1", project="demo", endpoint=EP, production=False,
                   probe_sets=["out_of_scope"], approval_id="ap1", planned_probes=1,
                   approved_probes=1, ceiling=50, records=[rec])
    write_json(store.paths.dir / "ai" / "adversarial" / "adv_1" / "run.json", run)


def _refusals(store: ProjectStore, lines: list[str] | None) -> None:
    path = store.paths.dir / "ai" / "adversarial" / "refusals.jsonl"
    path.parent.mkdir(parents=True, exist_ok=True)
    if lines is not None:
        path.write_text("".join(line + "\n" for line in lines), encoding="utf-8")


def _refusal_line(reason: str = "no signed approval") -> str:
    return ('{"target":"https://llm.test/chat","timestamp":"2026-10-11T00:00:00+00:00",'
            f'"requested_probes":15,"reason":"{reason}"}}')


def _build(store: ProjectStore, tmp_path: Path, checks: list[AiCheckEntry] | None = None,
           redactor: Redactor | None = None) -> Path:
    out = tmp_path / "out" / "ai.xlsx"
    return export_ai_report(store, out, checks=_checks() if checks is None else checks,
                            redactor=redactor or Redactor({}))


def _rows(path: Path, sheet: str) -> list[list[str]]:
    ws = load_workbook(path)[sheet]
    return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]


def _all_text(path: Path) -> str:
    wb = load_workbook(path)
    return "\n".join(str(c) for ws in wb for row in ws.iter_rows(values_only=True)
                     for c in row if c is not None)


def test_ai9_one_sheet_per_tier_in_tier_order_with_counts(tmp_path: Path) -> None:
    store = _store(tmp_path)
    _put(store, "k1", AiCheckKind.RESPONSE_SHAPE, _verdict(Result.PASS))
    out = _build(store, tmp_path)

    wb = load_workbook(out)
    assert wb.sheetnames == SHEETS
    assert [t.value for t in TIER_ORDER] == SHEETS[1:]
    summary = {r[0]: r for r in _rows(out, "Summary")}
    assert summary["static"][2] == "2"        # runnable
    assert summary["behavioural"][2] == "0"   # blocked, so 0 runnable
    assert summary["adversarial"][2] == "0"
    assert "2 runnable" in _rows(out, "static")[0][0]


def test_ai9_empty_tier_is_present_and_says_zero(tmp_path: Path) -> None:
    store = _store(tmp_path)
    _put(store, "k1", AiCheckKind.RESPONSE_SHAPE, _verdict(Result.PASS))
    out = _build(store, tmp_path, checks=_checks()[:2])  # nothing behavioural at all
    assert "behavioural" in load_workbook(out).sheetnames
    heading = _rows(out, "behavioural")[0][0]
    assert "0 runnable" in heading and "0 checks" in heading


def test_ai9_no_results_says_so(tmp_path: Path) -> None:
    text = _all_text(_build(_store(tmp_path), tmp_path, checks=[]))
    assert "No AI results are stored for this project" in text
    assert "all clear" not in text.lower()


def test_ai9_checks_without_stored_results_are_not_health(tmp_path: Path) -> None:
    out = _build(_store(tmp_path), tmp_path)
    assert "No AI results are stored for this project" in _all_text(out)
    assert [r[3] for r in _rows(out, "static")[2:] if r[1] == "response_shape"] == ["NOT RUN"]


@pytest.mark.parametrize("reason", list(BlockedReason))
def test_ai10_blocked_row_for_every_reason(tmp_path: Path, reason: BlockedReason) -> None:
    action = f"do the one thing that clears {reason.value}"
    entry = _entry(AiCheckKind.GROUND_TRUTH_ANSWER, Tier.BEHAVIOURAL, AiCheckMethod.JUDGED,
                   runnable=False, reason=reason, action=action)
    out = _build(_store(tmp_path), tmp_path, checks=[entry])

    [row] = [r for r in _rows(out, "behavioural") if r[1] == "ground_truth_answer"]
    assert row[3] == "BLOCKED"
    assert reason.value in row[-1] and action in row[-1]


def test_ai10_not_applicable_is_not_blocked_and_not_dropped(tmp_path: Path) -> None:
    na = AiCheckEntry(kind=AiCheckKind.TOOL_USE, tier=Tier.BEHAVIOURAL,
                      method=AiCheckMethod.JUDGED, applicable=False, runnable=False)
    out = _build(_store(tmp_path), tmp_path, checks=[na])
    assert [r[3] for r in _rows(out, "behavioural") if r[1] == "tool_use"] == ["NOT APPLICABLE"]


def test_ai11_every_failure_string_verbatim(tmp_path: Path) -> None:
    store = _store(tmp_path)
    fails = [Failure(criterion_id="AI-RS1", reason="Response  was NOT json; see {a} 'x'\nline2",
                     fix_hint="return an object"),
             Failure(criterion_id="AI-RS2", reason="missing keys: answer, sources")]
    _put(store, "k1", AiCheckKind.RESPONSE_SHAPE, _verdict(Result.FAIL, failures=fails))
    out = _build(store, tmp_path)

    rows = [r for r in _rows(out, "static") if r[1] == "response_shape" and r[3] == "FAIL"]
    assert [(r[7], r[8], r[9]) for r in rows] == [
        ("AI-RS1", fails[0].reason, "return an object"),
        ("AI-RS2", fails[1].reason, "")]


def test_ai11_inconclusive_note_verbatim_and_grader_labels(tmp_path: Path) -> None:
    store = _store(tmp_path)
    note = "The reference was too vague to decide."
    _put(store, "k1", AiCheckKind.RESPONSE_SHAPE, _verdict(Result.PASS))
    _put(store, "k2", AiCheckKind.GROUND_TRUTH_ANSWER,
         _verdict(Result.INCONCLUSIVE, note=note, provider="gemini"))
    out = _build(store, tmp_path, checks=[*_checks()[:1], _entry(
        AiCheckKind.GROUND_TRUTH_ANSWER, Tier.BEHAVIOURAL, AiCheckMethod.JUDGED)])

    beh = [r for r in _rows(out, "behavioural") if r[1] == "ground_truth_answer"]
    assert beh[0][3] == "INCONCLUSIVE" and beh[0][10] == note
    assert beh[0][5] == "judge: gemini"
    stat = [r for r in _rows(out, "static") if r[1] == "response_shape"]
    assert stat[0][5] == "computed (rule)"


def test_ai11_captured_but_not_judged_is_labelled(tmp_path: Path) -> None:
    store = _store(tmp_path)
    _put(store, "k1", AiCheckKind.RESPONSE_SHAPE, None)
    out = _build(store, tmp_path)
    rows = [r for r in _rows(out, "static") if r[1] == "response_shape"]
    assert [r[3] for r in rows] == ["CAPTURED, NOT JUDGED"]


def test_ai12_read_only_deterministic_and_no_model_or_network(tmp_path: Path) -> None:
    store = _store(tmp_path)
    _put(store, "k1", AiCheckKind.RESPONSE_SHAPE, _verdict(Result.PASS))
    _run(store)
    _refusals(store, [_refusal_line()])
    def tree() -> list[str]:
        return sorted(p.relative_to(tmp_path).as_posix() for p in tmp_path.rglob("*"))

    before = tree()

    with mock.patch("socket.socket.connect", side_effect=AssertionError("network")), \
            mock.patch.object(Provider, "judge", side_effect=AssertionError("model call")):
        a = _build(store, tmp_path)
        first = {s: _rows(a, s) for s in SHEETS}
        b = _build(store, tmp_path)
        second = {s: _rows(b, s) for s in SHEETS}

    assert first == second
    assert [p for p in tree() if p not in before] == ["out", "out/ai.xlsx"]


def test_ai12_stage_constructs_no_verdict() -> None:
    src = (ROOT / "src" / "autotester" / "stages" / "ai_report.py").read_text(encoding="utf-8")
    assert not re.search(r"Verdict\(", src)
    assert not re.search(r"\bgrade\(|\.judge\(|import socket|httpx|requests\.", src)


def test_ai13_two_refusals_two_rows_no_probe_text(tmp_path: Path) -> None:
    store = _store(tmp_path)
    _run(store)
    _refusals(store, [_refusal_line("no signed approval"), _refusal_line("over the ceiling")])
    out = _build(store, tmp_path)

    rows = _rows(out, "adversarial")
    refused = [r for r in rows if r[0] == "refused attempt"]
    assert [(r[1], r[2], r[3], r[4]) for r in refused] == [
        (EP, "2026-10-11T00:00:00+00:00", "15", "no signed approval"),
        (EP, "2026-10-11T00:00:00+00:00", "15", "over the ceiling")]
    assert "PROBE-TEXT" not in _all_text(out)


def test_ai13_missing_and_empty_logs_have_distinct_statements(tmp_path: Path) -> None:
    missing = _build(_store(tmp_path / "a"), tmp_path / "a")
    empty_store = _store(tmp_path / "b")
    _refusals(empty_store, [])
    empty = _build(empty_store, tmp_path / "b")

    m, e = _all_text(missing), _all_text(empty)
    assert "no refused attempts recorded" in m and "does not exist" in m
    assert "no refused attempts recorded" in e and "is empty" in e
    assert "does not exist" not in e


def test_ai13_unreadable_log_is_not_reported_as_empty(tmp_path: Path) -> None:
    store = _store(tmp_path)
    _refusals(store, [_refusal_line(), "{not json"])
    text = _all_text(_build(store, tmp_path))
    assert "could not be read" in text
    assert "no refused attempts recorded" not in text


def test_ai13_adversarial_run_rows_are_shown_unjudged(tmp_path: Path) -> None:
    store = _store(tmp_path)
    _run(store)
    rows = [r for r in _rows(_build(store, tmp_path), "adversarial") if r[0] == "run"]
    assert len(rows) == 1
    assert "adv_1" in rows[0][1] and "NOT JUDGED" in " ".join(rows[0])


def test_ai14_planted_secret_never_reaches_a_cell(tmp_path: Path) -> None:
    store = _store(tmp_path)
    fails = [Failure(criterion_id="AI-GT1", reason=f"answer leaked {SECRET}", fix_hint=None)]
    _put(store, "k1", AiCheckKind.RESPONSE_SHAPE, _verdict(Result.FAIL, failures=fails),
         response=f"here is {SECRET}")
    _refusals(store, [_refusal_line(f"bad {SECRET}")])
    out = _build(store, tmp_path, redactor=Redactor({"DEMO_KEY": SECRET}))

    assert SECRET not in _all_text(out)
    assert "[REDACTED]" in _all_text(out)
    env = tmp_path / ".env"
    env.write_text(f"DEMO_KEY={SECRET}\n", encoding="utf-8")
    code = subprocess.run(
        [sys.executable, str(ROOT / "scripts" / "check_no_secrets.py"), str(out)],
        cwd=tmp_path, capture_output=True, text=True, check=False)
    assert code.returncode == 0, code.stdout


def test_ai14_autosize_columns_is_the_shared_helper(tmp_path: Path) -> None:
    store = _store(tmp_path)
    _put(store, "k1", AiCheckKind.RESPONSE_SHAPE, _verdict(Result.PASS))
    with mock.patch.object(ai_report, "autosize_columns") as sizer:
        _build(store, tmp_path)
    assert sizer.call_count == len(SHEETS)


def test_ai14_control_characters_do_not_crash_the_workbook(tmp_path: Path) -> None:
    store = _store(tmp_path)
    fails = [Failure(criterion_id="AI-RS1", reason="bad\x07byte")]
    _put(store, "k1", AiCheckKind.RESPONSE_SHAPE, _verdict(Result.FAIL, failures=fails))
    assert "badbyte" in _all_text(_build(store, tmp_path))
