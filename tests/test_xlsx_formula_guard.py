"""AT-925: exporters must not write attacker-controlled text as live formulas."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from test_report_export_reason import RUN_ID, _seed

from autotester.core.excel import append_row, clean_cell
from autotester.schema.enums import Outcome, Result
from autotester.schema.verdict import Failure
from autotester.stages import report_export


@pytest.mark.parametrize("lead", ["=", "+", "-", "@", "\t", "\r"])
def test_clean_cell_neutralises_formula_leads(lead: str) -> None:
    assert clean_cell(f"{lead}cmd|' /C calc'!A0") == f"'{lead}cmd|' /C calc'!A0"


def test_clean_cell_leaves_plain_text_and_numbers() -> None:
    assert clean_cell("hello = world") == "hello = world"
    assert clean_cell(-5) == -5
    assert clean_cell(-1.5) == -1.5
    assert clean_cell(None) is None


def test_append_row_stores_text_not_formula_and_keeps_numbers() -> None:
    ws = Workbook().active
    append_row(ws, ["=1+1", -3])
    assert ws["A1"].data_type == "s"
    assert ws["B1"].value == -3
    assert ws["B1"].data_type == "n"


def test_export_excel_guards_formula_text_and_keeps_reason(tmp_path: Path) -> None:
    store = _seed(
        tmp_path, outcome=Outcome.COMPLETED, result=Result.FAIL,
        failures=[Failure(criterion_id="c1", reason="=cmd|' /C calc'!A0", fix_hint="@fix")],
        error="+boom",
    )
    out = report_export.export_excel("demo", RUN_ID, tmp_path / "o.xlsx", tmp_path)
    ws = load_workbook(out)["Run report"]
    assert store is not None
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            assert cell.data_type != "f"
    failures = ws.cell(row=2, column=10).value
    assert "reason" not in failures and "=cmd" in failures and "fix: @fix" in failures
