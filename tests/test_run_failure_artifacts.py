"""Offline production launcher/store/export proof for sanitized run-local failure bundles."""

from __future__ import annotations

import json
from types import SimpleNamespace

from openpyxl import load_workbook
from PIL import Image

from autotester.core import image_privacy
from autotester.schema.case import Case
from autotester.schema.enums import Action, CaseClass, CaseKind, EvidenceKind, Outcome, Result
from autotester.schema.failure_bundle import BundleIndex
from autotester.schema.flowspec import Step
from autotester.schema.project import Project, SecretRef
from autotester.schema.run import Evidence, RawResult
from autotester.schema.trace import LLMSpan, StageSpan
from autotester.schema.verdict import Verdict
from autotester.stages import run_artifacts, run_launcher
from autotester.stages.failure_bundle import load_bundle
from autotester.stages.report_export import export_excel, export_html
from autotester.store.project_store import ProjectStore

SECRET = "BundlePlantedCredential928"


def _launch(tmp_path, monkeypatch, *, paths=None, verdict_result=Result.FAIL):
    monkeypatch.setenv("AUTOTESTER_ROOT", str(tmp_path))
    store = ProjectStore("demo", tmp_path)
    project = Project(slug="demo", name="Demo", base_url="https://example.test",
                      secrets=[SecretRef(key="APP_PASSWORD", domains=["example.test"])])
    store.save_project(project)
    store.paths.env_file.write_text(f"APP_PASSWORD={SECRET}\n")
    case = Case(project="demo", flow_id="f", kind=CaseKind.BEST,
                case_class=CaseClass.AUTH_WRONG_CREDS, title="login",
                steps=[Step(order=1, action=Action.CLICK, target="button")])
    store.add_case(case)
    run_dir = store.paths.run_dir("run-A")
    run_dir.mkdir(parents=True)
    Image.new("RGB", (30, 20), "white").save(run_dir / "shot.png")
    result = RawResult(case_id=case.id, outcome=Outcome.ERRORED, error=SECRET,
                       evidence=[Evidence(kind=EvidenceKind.SCREENSHOT, path=p, step_order=1)
                                 for p in (paths or ["shot.png"])])
    verdict = Verdict(run_id="run-A", case_id=case.id, result=verdict_result, note=SECRET)
    original = {}

    def execute(*args):
        store.save_result("run-A", result)
        store.save_verdict("run-A", verdict)
        spans = [StageSpan(trace_id="run-A", stage="execute", status="failed", error=SECRET),
                 LLMSpan(trace_id="run-A", provider="fixture", role="judge", fed_id=case.id),
                 LLMSpan(trace_id="run-A", provider="other-case", role="judge", fed_id="other")]
        store.paths.run_trace("run-A").write_text(
            "\n".join(s.model_dump_json() for s in spans))
        original.update({p.name: p.read_bytes() for p in run_dir.glob("*.json")})
        return SimpleNamespace(n=1, bound_by="config")

    prepared = (store, project, [case], None, store.paths, None,
                SimpleNamespace(id="approval", max_actions=1, max_probes=1, wall_clock_s=30),
                [False], SimpleNamespace(runnable_in_tier=lambda tier: []))
    monkeypatch.setattr(run_launcher, "prepare_run", lambda *a: prepared)
    monkeypatch.setattr(run_launcher, "run_cases_traced", execute)
    run = run_launcher.launch_run("demo", run_id="run-A")
    return store, run, case, original


def _ocr(path):
    return ("level\tpage_num\tblock_num\tpar_num\tline_num\tconf\tleft\ttop\twidth\theight\ttext\n"
            "5\t1\t1\t1\t1\t99\t0\t6\t25\t1\tprivate\n")


def test_launcher_bundles_sanitized_copies_and_reports(tmp_path, monkeypatch):
    monkeypatch.setattr(image_privacy, "_local_ocr", _ocr)
    store, run, case, original = _launch(tmp_path, monkeypatch)
    index = BundleIndex.model_validate_json(
        (store.paths.bundles_dir(run.id) / "index.json").read_text())
    entry = index.cases[0]
    assert entry.status == "complete" and entry.images_included == 1
    assert entry.image_text_unavailable and entry.images_withheld == 0
    bundle = store.paths.bundles_dir(run.id) / entry.bundle_id
    manifest = load_bundle(bundle)
    copied = next(bundle / f.name for f in manifest.files if f.kind.value == "screenshot")
    with Image.open(copied) as shot:
        assert shot.getpixel((10, 7)) == (0, 0, 0)
    with Image.open(store.paths.run_dir(run.id) / "shot.png") as shot:
        assert shot.getpixel((10, 7)) == (255, 255, 255)
    assert all(SECRET not in p.read_text() for p in bundle.rglob("*")
               if p.is_file() and p.suffix in {".json", ".txt", ".jsonl"})
    assert "other-case" not in (bundle / "trace.jsonl").read_text()
    assert all((store.paths.run_dir(run.id) / n).read_bytes() == b for n, b in original.items())
    html = export_html("demo", run.id, tmp_path / "reports" / "report.html", tmp_path)
    assert "manifest.json'" in html.read_text()
    assert "image/text evidence unavailable" in html.read_text()
    excel = export_excel("demo", run.id, tmp_path / "reports" / "report.xlsx", tmp_path)
    cell = load_workbook(excel).active.cell(2, 12)
    assert (excel.parent / cell.hyperlink.target).resolve() == bundle / "manifest.json"
    assert len(store.load_results(run.id)) == 1  # nested index never enters RawResult glob
    assert store.load_run(run.id).case_ids == [case.id]


def test_bundle_exception_preserves_functional_records(tmp_path, monkeypatch, caplog):
    def fail(**kwargs):
        raise RuntimeError(SECRET)

    monkeypatch.setattr(run_artifacts, "build_failure_bundle", fail)
    store, run, case, original = _launch(tmp_path, monkeypatch)
    assert all((store.paths.run_dir(run.id) / n).read_bytes() == b for n, b in original.items())
    path = store.paths.bundles_dir(run.id) / "index.json"
    index = BundleIndex.model_validate_json(path.read_text())
    assert index.cases[0].status == "unavailable" and index.cases[0].diagnostic == "assembly_failed"
    assert SECRET not in path.read_text()
    assert run_artifacts.bundle_report_detail(store, run.id, case.id)[0] is None
    assert store.load_verdicts(run.id)[0].result is Result.FAIL
    monkeypatch.setattr(run_artifacts, "write_json", lambda *a: fail())
    store, run, _, original = _launch(tmp_path / "index-failure", monkeypatch)
    assert all((store.paths.run_dir(run.id) / n).read_bytes() == b for n, b in original.items())
    assert "failure_bundle_index_unavailable" in caplog.text and SECRET not in caplog.text


def test_containment_and_uncertain_images_are_explicit(tmp_path, monkeypatch):
    outside = tmp_path / "outside.png"
    Image.new("RGB", (30, 20), "white").save(outside)
    inspected = []

    def uncertain(path):
        inspected.append(path)
        return _ocr(path).replace("99\t", "101\t")  # impossible score: still fail closed

    monkeypatch.setattr(image_privacy, "_local_ocr", uncertain)
    store, run, case, _ = _launch(tmp_path, monkeypatch, paths=[str(outside), "shot.png"],
                                 verdict_result=Result.BLOCKED)
    index = BundleIndex.model_validate_json(
        (store.paths.bundles_dir(run.id) / "index.json").read_text())
    entry = index.cases[0]
    assert entry.status == "complete" and entry.images_included == 0
    assert entry.images_withheld == 2 and entry.image_text_unavailable
    assert entry.image_reasons == ["artifact_outside_run", "uncertain_ocr"]
    assert len(inspected) == 1
    assert "unavailable" in run_artifacts.bundle_report_detail(store, run.id, case.id)[1]
    result = store.load_results(run.id)[0]
    result.evidence.append(Evidence(kind=EvidenceKind.SCREENSHOT, path="shot.png",
                                    step_order=1, masked=False))
    store.save_result(run.id, result)
    run_artifacts.assemble_run_bundles(store, run)
    entry = BundleIndex.model_validate_json(
        (store.paths.bundles_dir(run.id) / "index.json").read_text()).cases[0]
    assert "unmasked_or_unattributed_image" in entry.image_reasons and len(inspected) == 2
    run.case_ids = []
    run_artifacts.assemble_run_bundles(store, run)
    assert run_artifacts.bundle_report_detail(store, run.id, case.id)[0] is None
    run.case_ids = [case.id]
    store.paths.run_trace(run.id).write_text(
        StageSpan(trace_id="run-B", stage="execute", status="failed").model_dump_json())
    run_artifacts.assemble_run_bundles(store, run)
    assert run_artifacts.bundle_report_detail(store, run.id, case.id)[0] is None


def test_report_lookup_rejects_partial_corrupt_and_wrong_lineage(tmp_path, monkeypatch):
    monkeypatch.setattr(image_privacy, "_local_ocr", _ocr)
    store, run, case, _ = _launch(tmp_path, monkeypatch, verdict_result=Result.INCONCLUSIVE)
    index_path = store.paths.bundles_dir(run.id) / "index.json"
    index_text = index_path.read_text()
    row = json.loads(index_text)
    bundle = store.paths.bundles_dir(run.id) / row["cases"][0]["bundle_id"]
    manifest_path = bundle / "manifest.json"
    manifest_text = manifest_path.read_text()
    assert run_artifacts.bundle_report_detail(store, run.id, case.id)[0]
    for key, value in [("run_id", "run-B"), ("case_id", "other"), ("bundle_id", "fb_other")]:
        manifest = json.loads(manifest_text)
        manifest[key] = value
        manifest_path.write_text(json.dumps(manifest))
        assert run_artifacts.bundle_report_detail(store, run.id, case.id)[0] is None
        manifest_path.write_text(manifest_text)
    partial = bundle.with_name(bundle.name + ".partial")
    bundle.rename(partial)
    assert run_artifacts.bundle_report_detail(store, run.id, case.id)[0] is None
    partial.rename(bundle)
    error_text = (bundle / "error.txt").read_text()
    (bundle / "error.txt").write_text("corrupt")
    assert run_artifacts.bundle_report_detail(store, run.id, case.id)[0] is None
    html = export_html("demo", run.id, tmp_path / "report.html", tmp_path).read_text()
    assert "Failure bundle unavailable" in html and "manifest.json'" not in html
    (bundle / "error.txt").write_text(error_text)
    assert run_artifacts.bundle_report_detail(store, run.id, case.id)[0]
    row["cases"][0]["status"] = "unavailable"
    index_path.write_text(json.dumps(row))
    assert run_artifacts.bundle_report_detail(store, run.id, case.id)[0] is None
    row["cases"][0]["status"] = "complete"
    row["run_id"] = "run-B"
    index_path.write_text(json.dumps(row))
    assert run_artifacts.bundle_report_detail(store, run.id, case.id)[0] is None
    store, run, case, _ = _launch(tmp_path / "passed", monkeypatch, verdict_result=Result.PASS)
    assert BundleIndex.model_validate_json(
        (store.paths.bundles_dir(run.id) / "index.json").read_text()).cases == []
