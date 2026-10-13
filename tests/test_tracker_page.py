"""T-199 tracker page and Tracker-layout export (TB7 in part, TB15). Contract: team-bug-loop.md."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

import pytest
from bug_loop_fixture import event, seed
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from autotester.core.paths import ProjectPaths
from autotester.schema.project import Project, SecretRef
from autotester.schema.tracked_bug import BugStatus, EventKind
from autotester.stages import sheet_plan, tracker_export
from autotester.stages.sheet_write import LiveWriteRefused
from autotester.store.bug_store import BugStore
from autotester.store.project_store import ProjectStore
from autotester.ui.app import app
from autotester.ui.routes_tracker import ROUTE_POLICY, router

K = EventKind
SECRET = "hunter2-PLANTED-9f3a"
HEADERS = ["Bug ID", "Title", "Severity", "Screen", "Status", "Assigned To", "First Seen",
           "Last Seen", "Count", "Evidence"]


@pytest.fixture
def bugs(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> BugStore:
    monkeypatch.setenv("AUTOTESTER_ROOT", str(tmp_path))
    ProjectStore("demo", tmp_path).save_project(Project(
        slug="demo", name="Demo", base_url="https://demo.test/", allowed_domains=["demo.test"],
        secrets=[SecretRef(key="DEMO_PASSWORD", domains=["demo.test"])]))
    (tmp_path / ".env").write_text(f"DEMO_PASSWORD={SECRET}\n", encoding="utf-8")
    store = BugStore(ProjectPaths("demo", tmp_path).dir)
    first = seed("case_aaaaaaaaaaaa", "C1", f"Checkout fails for {SECRET}")
    second = seed("case_bbbbbbbbbbbb", "C2", "=HYPERLINK(\"http://evil\") login")
    other = seed("case_cccccccccccc", "C3", "Footer typo")
    store.apply_all([
        event(K.FAIL, 1, seed_obj=first), event(K.FAIL, 2, seed_obj=first),
        event(K.FAIL, 3, seed_obj=second),
        event(K.FAIL, 4, seed_obj=other),
        event(K.ASSIGN, 5, key=other.bug_key, assignee="priya"),
    ])
    return store


def _client() -> TestClient:
    return TestClient(app)


def test_list_shows_every_bug_with_its_fields_and_evidence(bugs: BugStore) -> None:
    html = _client().get("/projects/demo/tracker").text
    assert "3 of 3 bug(s)" in html and "Footer typo" in html and "Checkout" in html
    assert "/projects/demo/runs/run2" in html and "priya" in html
    assert SECRET not in html


def test_status_and_severity_filters(bugs: BugStore) -> None:
    key = next(b.bug_key for b in bugs.bugs().values() if b.title == "Footer typo")
    html = _client().get("/projects/demo/tracker?status=fixed").text
    assert "0 of 3" in html and "No bugs match" in html
    html = _client().get("/projects/demo/tracker?status=open&severity=S2").text
    assert "3 of 3" in html and key in html
    assert _client().get("/projects/demo/tracker?status=bogus").status_code == 400
    assert _client().get("/projects/demo/tracker?severity=S9").status_code == 400


def test_detail_shows_event_history_and_unknown_bug_is_404(bugs: BugStore) -> None:
    key = next(b.bug_key for b in bugs.bugs().values() if b.title == "Footer typo")
    html = _client().get(f"/projects/demo/tracker/{key}").text
    assert "Footer typo" in html and "assign" in html and "fail" in html and "priya" in html
    assert _client().get("/projects/demo/tracker/nope").status_code == 404
    assert _client().get("/projects/ghost/tracker").status_code == 404


def test_empty_store_lists_nothing(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("AUTOTESTER_ROOT", str(tmp_path))
    ProjectStore("demo", tmp_path).save_project(Project(
        slug="demo", name="Demo", base_url="https://demo.test/", allowed_domains=["demo.test"]))
    assert "No bugs match" in _client().get("/projects/demo/tracker").text


def _csv_rows(bugs: BugStore) -> list[list[str]]:
    response = _client().get("/projects/demo/tracker/export.csv")
    assert response.headers["content-type"].startswith("text/csv")
    return list(csv.reader(io.StringIO(response.content.decode("utf-8"))))


def test_csv_follows_the_layout_and_matches_the_bugs(bugs: BugStore) -> None:
    rows = _csv_rows(bugs)
    assert rows[0] == HEADERS == tracker_export.header_row()
    by_title = {r[1]: r for r in rows[1:]}
    assert len(rows) == 4
    footer = by_title["Footer typo"]
    bug = next(b for b in bugs.bugs().values() if b.title == "Footer typo")
    assert footer[0] == bug.bug_key and footer[2] == "S2" and footer[4] == "open"
    assert footer[5] == "priya" and footer[8] == "1" and footer[9] == "/projects/demo/runs/run4"
    assert rows[1][1].startswith("Checkout fails")  # first seen first


def test_xlsx_has_the_same_header_and_cells_as_the_csv(bugs: BugStore) -> None:
    response = _client().get("/projects/demo/tracker/export.xlsx")
    sheet = load_workbook(io.BytesIO(response.content)).active
    cells = [[str(c) if c is not None else "" for c in row]
             for row in sheet.iter_rows(values_only=True)]
    assert sheet.title == "Tracker" and cells == _csv_rows(bugs)


def test_planted_secret_and_formula_are_neutralised_in_both_files(bugs: BugStore) -> None:
    csv_bytes = _client().get("/projects/demo/tracker/export.csv").content
    xlsx_bytes = _client().get("/projects/demo/tracker/export.xlsx").content
    assert SECRET.encode() not in csv_bytes
    sheet = load_workbook(io.BytesIO(xlsx_bytes)).active
    flat = " ".join(str(c) for row in sheet.iter_rows(values_only=True) for c in row)
    assert SECRET not in flat
    titles = [row[1] for row in csv.reader(io.StringIO(csv_bytes.decode()))]
    assert "'=HYPERLINK(\"http://evil\") login" in titles


def test_re_export_gives_identical_bytes(bugs: BugStore) -> None:
    for path in ("export.csv", "export.xlsx"):
        first = _client().get(f"/projects/demo/tracker/{path}").content
        second = _client().get(f"/projects/demo/tracker/{path}").content
        assert first == second


def test_export_of_unknown_project_is_not_found(bugs: BugStore) -> None:
    assert _client().get("/projects/ghost/tracker/export.csv").status_code == 404
    assert _client().get("/projects/ghost/tracker/export.xlsx").status_code == 404


def test_layout_reuses_the_sheet_plan_mapping_and_reordering_it_breaks_the_export() -> None:
    assert tracker_export.bug_values is sheet_plan.bug_values
    assert tracker_export.safe_cell is sheet_plan.safe_cell
    assert list(tracker_export.TRACKER_COLUMNS.columns) == [
        "bug_key", "title", "severity", "screen", "status", "assignee", "first_seen",
        "last_seen", "count"]


def test_export_module_never_touches_a_sheet_client_or_the_network() -> None:
    src = Path(tracker_export.__file__).read_text(encoding="utf-8")
    code = re.sub(r'""".*?"""', "", src, flags=re.S)
    assert not re.search(r"SheetClient|requests|httpx|urllib|socket|gspread|googleapiclient", code)
    assert issubclass(LiveWriteRefused, Exception)


def test_every_tracker_route_declares_report_view_policy() -> None:
    served = {f"{m} {r.path}" for r in router.routes for m in r.methods}
    assert served and all(r.path.endswith(("tracker", "csv", "xlsx", "{bug_key}"))
                          for r in router.routes)
    assert served == set(ROUTE_POLICY)
    assert {v[0] for v in ROUTE_POLICY.values()} == {"report.view"}


def test_a_fixed_bug_filters_as_fixed(bugs: BugStore) -> None:
    key = next(b.bug_key for b in bugs.bugs().values() if b.title == "Footer typo")
    bugs.apply_all([event(K.PASS, 6, key=key, minutes=60), event(K.PASS, 7, key=key, minutes=70)])
    assert bugs.get(key).status is BugStatus.FIXED
    html = _client().get("/projects/demo/tracker?status=fixed").text
    assert "1 of 3" in html and "Footer typo" in html
